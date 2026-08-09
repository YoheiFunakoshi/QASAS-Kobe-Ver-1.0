import csv
from pathlib import Path
import unittest

from openpyxl import Workbook

from qasas.loaders import detect_sample_format, load_cpm, load_database, load_rg
from tests.support import test_path


class LoaderTests(unittest.TestCase):
    def test_cpm_loader_aggregates_and_recalculates_frequency(self):
        path = test_path("cpm_sample.csv")
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=["Vseg", "Jseg", "CDR3", "Counts", "Frequency(%)"])
            writer.writeheader()
            writer.writerow({"Vseg": "IGHV3-23*01", "Jseg": "IGHJ4*02", "CDR3": "CARDRW", "Counts": 3})
            writer.writerow({"Vseg": "IGHV3-23*02", "Jseg": "IGHJ4*01", "CDR3": "CARDRW", "Counts": 2})
            writer.writerow({"Vseg": "IGHV1-2*01", "Jseg": "IGHJ6*01", "CDR3": "CQQQW", "Counts": 5})
        self.assertEqual(detect_sample_format(path), "CPM")
        sample = load_cpm(path)
        self.assertEqual(len(sample.clones), 2)
        self.assertEqual(sample.total_reads, 10)
        self.assertEqual(sample.clones[0].frequency_percent, 50.0)
        self.assertEqual(sum(clone.frequency_percent for clone in sample.clones), 100.0)

    def test_rg_loader_uses_back_data_and_in_frame_only(self):
        path = test_path("rg_sample.xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Back_data"
        sheet["A1"] = "Sample ID"
        sheet["C1"] = "RG001"
        sheet["A2"] = "In-frame reads"
        sheet["C2"] = 12
        sheet.cell(10, 7, "IGHV3-23*01")
        sheet.cell(10, 11, "IGHJ4*02")
        sheet.cell(10, 15, "CARDRW")
        sheet.cell(10, 16, "in-frame")
        sheet.cell(10, 17, 7)
        sheet.cell(11, 7, "IGHV3-23*02")
        sheet.cell(11, 11, "IGHJ4*01")
        sheet.cell(11, 15, "CARDRW")
        sheet.cell(11, 16, "in-frame")
        sheet.cell(11, 17, 5)
        sheet.cell(12, 7, "IGHV1-2*01")
        sheet.cell(12, 11, "IGHJ6*01")
        sheet.cell(12, 15, "COUTW")
        sheet.cell(12, 16, "out-of-frame")
        sheet.cell(12, 17, 99)
        workbook.save(path)
        sample = load_rg(path)
        self.assertEqual(sample.sample_id, "RG001")
        self.assertEqual(len(sample.clones), 1)
        self.assertEqual(sample.total_reads, 12)
        self.assertEqual(sample.listed_reads, 12)
        self.assertEqual(sample.source_rows, 3)
        self.assertEqual(sample.skipped_rows, 1)

    def test_database_loader_deduplicates_v_j_cdr3_and_merges_annotations(self):
        path = test_path("database.csv")
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=["Name", "Heavy V Gene", "Heavy J Gene", "CDRH3", "Binds to"],
            )
            writer.writeheader()
            writer.writerow(
                {"Name": "Ab1", "Heavy V Gene": "IGHV3-23*01 (Human)", "Heavy J Gene": "IGHJ4*02", "CDRH3": "ARDR", "Binds to": "A"}
            )
            writer.writerow(
                {"Name": "Ab2", "Heavy V Gene": "IGHV3-23*02", "Heavy J Gene": "IGHJ4*01", "CDRH3": "ARDR", "Binds to": "B"}
            )
        database = load_database(path)
        self.assertEqual(database.usable_rows, 2)
        self.assertEqual(len(database.entries), 1)
        self.assertEqual(database.entries[0].annotations["Name"], ("Ab1", "Ab2"))


if __name__ == "__main__":
    unittest.main()
