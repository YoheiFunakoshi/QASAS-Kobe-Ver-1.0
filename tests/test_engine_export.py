from pathlib import Path
import unittest

from openpyxl import load_workbook

from qasas.engine import analyse
from qasas.export import export_result_xlsx
from qasas.models import DatabaseData, DatabaseEntry, SampleClone, SampleData
from tests.support import test_path


class EngineExportTests(unittest.TestCase):
    def _sample(self) -> SampleData:
        clones = (
            SampleClone(("IGHV1-2",), ("IGHJ4",), "ABCDE", 10, 50.0),
            SampleClone(("IGHV1-2",), ("IGHJ4",), "ABCXE", 5, 25.0),
            SampleClone(("IGHV1-2",), ("IGHJ4",), "ABXYE", 3, 15.0),
            SampleClone(("IGHV3-23",), ("IGHJ6",), "ABCDE", 2, 10.0),
        )
        return SampleData(Path("sample.csv"), "CPM", "S1", clones, 20, 4, 4, 0)

    def _database(self) -> DatabaseData:
        entries = (
            DatabaseEntry("IGHV1-2", "IGHJ4", "ABCDE", {"Name": ("KnownAb",)}),
            DatabaseEntry("IGHV3-23", "IGHJ4", "ABCDE", {"Name": ("WrongJ",)}),
        )
        return DatabaseData(
            Path("db.csv"), entries, 2, 2, 0, ("Name",), "Heavy V Gene", "Heavy J Gene", "CDRH3"
        )

    def test_analysis_keeps_exact_distance_classes_separate(self):
        result = analyse(self._sample(), self._database())
        self.assertEqual([item.unique_clones for item in result.exact_summaries], [1, 1, 1])
        self.assertEqual([item.total_reads for item in result.exact_summaries], [10, 5, 3])
        self.assertEqual([item.unique_clones for item in result.cumulative_summaries], [1, 2, 3])
        self.assertEqual(result.matched_clone_count, 3)

    def test_export_contains_required_sheets(self):
        result = analyse(self._sample(), self._database())
        output = export_result_xlsx(result, test_path("result.xlsx"))
        workbook = load_workbook(output, read_only=True, data_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["Summary", "Input QC", "Matched Clones"])
            self.assertEqual(workbook["Matched Clones"]["A2"].value, "LV0")
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
